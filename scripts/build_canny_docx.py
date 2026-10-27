import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_preprocessing_docx as base

ROOT = base.ROOT
base.SOURCE = ROOT / "Bao_cao_Phan_3_Canny.md"
base.OUTPUT = ROOT / "Bao_cao_Phan_3_Canny.docx"


def front_matter(document):
    info = document.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_after = Pt(16)
    base.apply_font(info.add_run("Module Canny"), size=Pt(11.5), bold=True, color=base.ACCENT)


base.add_front_matter = front_matter
path = base.build()

doc = Document(path)
for run in doc.sections[0].header.paragraphs[0].runs:
    run.text = "PHẦN 3 · CANNY"
doc.core_properties.title = "Báo cáo Phần 3 - Canny"
doc.core_properties.subject = "Canny phát hiện cạnh cho biên đường"
doc.core_properties.author = ""
doc.core_properties.keywords = "Canny, Sobel, hysteresis, threshold, road boundary"
doc.save(path)
print(path)
