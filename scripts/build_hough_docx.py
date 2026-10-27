import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_preprocessing_docx as base

ROOT = base.ROOT
base.SOURCE = ROOT / "baocao" / "Bao_cao_Phan_4_Hough.md"
base.OUTPUT = ROOT / "baocao" / "Bao_cao_Phan_4_Hough.docx"


def front_matter(document: Document) -> None:
    info = document.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_after = Pt(16)
    base.apply_font(info.add_run("Module Biến đổi Hough"), size=Pt(11.5), bold=True, color=base.ACCENT)


base.add_front_matter = front_matter
path = base.build()

doc = Document(path)
for run in doc.sections[0].header.paragraphs[0].runs:
    run.text = "PHẦN 4 · BIẾN ĐỔI HOUGH"
doc.core_properties.title = "Báo cáo Phần 4 - Biến đổi Hough"
doc.core_properties.subject = "Biến đổi Hough và hậu xử lý phát hiện hai biên đường"
doc.core_properties.author = ""
doc.core_properties.keywords = "Hough transform, Probabilistic Hough, slope filtering, boundary fitting, road boundary"
doc.save(path)
print(f"Da tao bao cao DOCX: {path}")
