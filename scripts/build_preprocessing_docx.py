from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Bao_cao_Phan_2_Tien_xu_ly_HSV.md"
OUTPUT = ROOT / "Bao_cao_Phan_2_Tien_xu_ly_HSV.docx"

BODY_FONT = "Calibri"
BODY_SIZE = Pt(12.5)
TEXT_COLOR = RGBColor(0, 0, 0)
ACCENT = RGBColor(31, 78, 121)
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_table_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))


def apply_font(run, *, size=None, bold=None, italic=None, color=None, name=BODY_FONT) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


INLINE_PATTERN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")


def add_inline_runs(paragraph, text: str) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            apply_font(paragraph.add_run(text[position : match.start()]), size=BODY_SIZE, color=TEXT_COLOR)
        token = match.group(0)
        if token.startswith("**"):
            apply_font(paragraph.add_run(token[2:-2]), size=BODY_SIZE, bold=True, color=TEXT_COLOR)
        elif token.startswith("`"):
            apply_font(
                paragraph.add_run(token[1:-1]),
                size=Pt(11.5),
                color=TEXT_COLOR,
                name="Consolas",
            )
        else:
            apply_font(paragraph.add_run(token[1:-1]), size=BODY_SIZE, italic=True, color=TEXT_COLOR)
        position = match.end()
    if position < len(text):
        apply_font(paragraph.add_run(text[position:]), size=BODY_SIZE, color=TEXT_COLOR)


def style_paragraph(paragraph, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=Pt(6), line=1.18) -> None:
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_after = after
    fmt.line_spacing = line
    fmt.widow_control = True


def add_body_paragraph(document: Document, text: str):
    paragraph = document.add_paragraph()
    add_inline_runs(paragraph, text)
    style_paragraph(paragraph)
    return paragraph


def add_list_item(document: Document, text: str, *, ordered: bool = False):
    style_name = "List Number" if ordered else "List Bullet"
    paragraph = document.add_paragraph(style=style_name)
    add_inline_runs(paragraph, text)
    style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.LEFT, after=Pt(3), line=1.12)
    return paragraph


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = len(rows[0])
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        keep_table_row(row)
        if row_index == 0:
            repeat_table_header(row)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline_runs(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
                if row_index == 0:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(document: Document, relative_path: str, caption: str | None) -> None:
    image_path = (ROOT / relative_path).resolve()
    if not image_path.is_file():
        raise FileNotFoundError(image_path)
    if image_path.name == "figure_3_before_after_gaussian.png":
        document.add_page_break()
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.space_before = Pt(4)
    picture_paragraph.paragraph_format.space_after = Pt(3)
    picture_paragraph.paragraph_format.keep_together = True
    picture_paragraph.add_run().add_picture(str(image_path), width=Cm(14.2))

    if caption:
        caption_paragraph = document.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph.paragraph_format.space_after = Pt(8)
        caption_paragraph.paragraph_format.keep_together = True
        run = caption_paragraph.add_run(caption)
        apply_font(run, size=Pt(10.5), italic=True, color=TEXT_COLOR)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = BODY_SIZE
    normal.font.color.rgb = TEXT_COLOR

    title = document.styles["Title"]
    title.font.name = BODY_FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = TEXT_COLOR
    title.paragraph_format.space_after = Pt(8)

    for style_name, size, before, after in (
        ("Heading 1", Pt(15.5), Pt(12), Pt(5)),
        ("Heading 2", Pt(13.5), Pt(9), Pt(4)),
    ):
        style = document.styles[style_name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = size
        style.font.bold = True
        style.font.color.rgb = TEXT_COLOR
        style.paragraph_format.space_before = before
        style.paragraph_format.space_after = after
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    header_run = header.add_run("PHẦN 2 · TIỀN XỬ LÝ HSV")
    apply_font(header_run, size=Pt(9), bold=True, color=ACCENT)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Trang ")
    apply_font(footer_run, size=Pt(9), color=RGBColor(89, 89, 89))
    add_field(footer, "PAGE")


def parse_markdown(document: Document, text: str) -> None:
    lines = text.splitlines()
    index = 0
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        if paragraph_buffer:
            add_body_paragraph(document, " ".join(part.strip() for part in paragraph_buffer))
            paragraph_buffer.clear()

    while index < len(lines):
        line = lines[index].rstrip()

        if not line:
            flush_paragraph()
            index += 1
            continue

        if line.startswith("# "):
            flush_paragraph()
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(line[2:].strip())
            apply_font(run, size=Pt(20), bold=True, color=TEXT_COLOR)
            index += 1
            continue

        if line.startswith("## "):
            flush_paragraph()
            document.add_heading(line[3:].strip(), level=1)
            index += 1
            continue

        if line.startswith("### "):
            flush_paragraph()
            document.add_heading(line[4:].strip(), level=2)
            index += 1
            continue

        if line.startswith("!["):
            flush_paragraph()
            image_match = re.match(r"!\[[^]]*\]\(([^)]+)\)", line)
            if not image_match:
                raise ValueError(f"Dòng ảnh không hợp lệ: {line}")
            caption = None
            if index + 1 < len(lines):
                next_line = lines[index + 1].strip()
                if next_line.startswith("*") and next_line.endswith("*"):
                    caption = next_line[1:-1]
                    index += 1
            add_figure(document, image_match.group(1), caption)
            index += 1
            continue

        if line.startswith("|"):
            flush_paragraph()
            table_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows: list[list[str]] = []
            for table_index, table_line in enumerate(table_lines):
                values = [cell.strip() for cell in table_line.strip("|").split("|")]
                if table_index == 1 and all(re.fullmatch(r":?-{3,}:?", value) for value in values):
                    continue
                rows.append(values)
            add_markdown_table(document, rows)
            continue

        if line.startswith("```"):
            flush_paragraph()
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].startswith("```"):
                code_lines.append(lines[index])
                index += 1
            code_paragraph = document.add_paragraph()
            code_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            code_paragraph.paragraph_format.left_indent = Cm(0.35)
            code_paragraph.paragraph_format.right_indent = Cm(0.35)
            code_paragraph.paragraph_format.space_before = Pt(3)
            code_paragraph.paragraph_format.space_after = Pt(7)
            code_paragraph.paragraph_format.keep_together = True
            set_cell_like_shading(code_paragraph, LIGHT_GRAY)
            run = code_paragraph.add_run("\n".join(code_lines))
            apply_font(run, size=Pt(9.5), color=TEXT_COLOR, name="Consolas")
            index += 1
            continue

        ordered_match = re.match(r"^\d+\.\s+(.*)$", line)
        if ordered_match:
            flush_paragraph()
            add_list_item(document, ordered_match.group(1), ordered=True)
            index += 1
            continue

        if line.startswith("- "):
            flush_paragraph()
            add_list_item(document, line[2:].strip(), ordered=False)
            index += 1
            continue

        if line.startswith("*") and line.endswith("*") and not line.startswith("**"):
            flush_paragraph()
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(10)
            run = paragraph.add_run(line[1:-1])
            apply_font(run, size=Pt(12.5), italic=True, color=RGBColor(89, 89, 89))
            index += 1
            continue

        paragraph_buffer.append(line)
        index += 1

    flush_paragraph()


def set_cell_like_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_front_matter(document: Document) -> None:
    info = document.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_after = Pt(16)
    run = info.add_run("Người thực hiện: Thế · Module tiền xử lý ảnh")
    apply_font(run, size=Pt(11.5), bold=True, color=ACCENT)

    box = document.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    box.autofit = False
    cell = box.cell(0, 0)
    cell.width = Cm(15.2)
    set_cell_shading(cell, "EAF2F8")
    set_cell_margins(cell, top=160, start=180, bottom=160, end=180)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(
        "Phạm vi bàn giao: resize → RGB sang HSV → lọc trắng/vàng → Gaussian → grayscale. "
        "Không triển khai Canny, ROI hoặc Hough trong module này."
    )
    apply_font(run, size=Pt(11.5), bold=True, color=TEXT_COLOR)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build() -> Path:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)

    document = Document()
    configure_document(document)

    markdown = SOURCE.read_text(encoding="utf-8")
    lines = markdown.splitlines()
    if not lines or not lines[0].startswith("# "):
        raise ValueError("Báo cáo Markdown phải bắt đầu bằng tiêu đề cấp 1")

    parse_markdown(document, markdown)

    title_paragraph = document.paragraphs[0]
    # python-docx không có API chèn theo chỉ mục; di chuyển XML của phần thông tin lên sau phụ đề.
    info_document = Document()
    configure_document(info_document)
    add_front_matter(info_document)
    anchor = document.paragraphs[1]._p if len(document.paragraphs) > 1 else title_paragraph._p
    for element in reversed(list(info_document._body._body)[:-1]):
        anchor.addnext(element)

    properties = document.core_properties
    properties.title = "Báo cáo Phần 2 - Tiền xử lý ảnh bằng HSV"
    properties.subject = "Tiền xử lý dữ liệu đầu vào cho Canny trong bài toán phát hiện biên làn"
    properties.author = "Thế"
    properties.keywords = "HSV, lane boundary, color filtering, Gaussian, preprocessing"

    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
