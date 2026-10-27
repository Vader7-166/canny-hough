import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
ROOT_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_preprocessing_docx as base

base.ROOT = ROOT_DIR
base.SOURCE = ROOT_DIR / "baocao" / "hough_road_boundary_tom_tat.md"
base.OUTPUT = ROOT_DIR / "baocao" / "hough_road_boundary_tom_tat.docx"


def front_matter(document: Document) -> None:
    info = document.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_after = Pt(16)
    base.apply_font(info.add_run("Tài liệu tóm tắt lý thuyết và thực nghiệm · Module Hough"), size=Pt(11.5), bold=True, color=base.ACCENT)


base.add_front_matter = front_matter
path = base.build()

doc = Document(path)
for run in doc.sections[0].header.paragraphs[0].runs:
    run.text = "TÓM TẮT · PHẦN 4 HOUGH"
doc.core_properties.title = "Phần Hough trong phát hiện biên đường"
doc.core_properties.subject = "Tài liệu tóm tắt cho người phụ trách Hough"
doc.save(path)
print(f"Da tao: {path}")
